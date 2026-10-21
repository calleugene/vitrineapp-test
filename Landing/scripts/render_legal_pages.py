from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "legal"
SOURCES = [
    (
        Path(r"C:\Users\user\Desktop\Политика обработки персональных данных.docx"),
        "politika-obrabotki-personalnyh-dannyh",
    ),
    (
        Path(r"C:\Users\user\Desktop\Согласие на обработку персональных данных.docx"),
        "soglasie-na-obrabotku-personalnyh-dannyh",
    ),
    (
        Path(r"C:\Users\user\Desktop\Согласие на получение рекламных рассылок.docx"),
        "soglasie-na-reklamnye-rassylki",
    ),
]


def page(title: str, body: str) -> str:
    return f"""<!doctype html>
<html lang=\"ru\">
<head>
  <meta charset=\"UTF-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>{escape(title)} — Витрина</title>
  <style>
    @font-face {{ font-family: Inter; src: url('../fonts/inter-cyrillic-wght-normal.woff2') format('woff2'); font-weight: 100 900; }}
    @font-face {{ font-family: Cormorant; src: url('../fonts/cormorant-garamond-cyrillic-wght-normal.woff2') format('woff2'); font-weight: 300 700; }}
    :root {{ color: #171512; background: #faf8f4; font-family: Inter, Arial, sans-serif; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; }}
    main {{ width: min(920px, calc(100% - 40px)); margin: 0 auto; padding: 58px 0 86px; }}
    a {{ color: inherit; }}
    .top {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 1px solid #e4ded6; padding-bottom: 28px; margin-bottom: 46px; font-size: 13px; }}
    .brand {{ letter-spacing: .14em; text-decoration: none; font-weight: 650; }}
    h1 {{ font-family: Cormorant, Georgia, serif; font-size: clamp(42px, 6vw, 66px); line-height: .98; font-weight: 500; letter-spacing: -.03em; margin: 0 0 42px; max-width: 760px; }}
    h2 {{ font-family: Cormorant, Georgia, serif; font-size: 31px; line-height: 1.1; font-weight: 500; margin: 46px 0 16px; }}
    p {{ max-width: 780px; margin: 0 0 17px; color: #403c37; font-size: 15px; line-height: 1.75; }}
    table {{ width: 100%; margin: 28px 0; border-collapse: collapse; font-size: 13px; line-height: 1.55; }}
    th, td {{ padding: 14px; text-align: left; vertical-align: top; border: 1px solid #e4ded6; }}
    th {{ background: #f4efe8; font-weight: 600; }}
    @media (max-width: 600px) {{ main {{ width: min(100% - 30px, 920px); padding-top: 28px; }} .top {{ margin-bottom: 34px; }} p {{ font-size: 14px; }} table {{ display: block; overflow-x: auto; }} }}
  </style>
</head>
<body>
  <main>
    <nav class=\"top\" aria-label=\"Навигация\"><a class=\"brand\" href=\"../\">ВИТРИНА</a><a href=\"../\">Вернуться на сайт</a></nav>
    <article>
      <h1>{escape(title)}</h1>
      {body}
    </article>
  </main>
</body>
</html>"""


def is_heading(style: str, text: str) -> bool:
    return style.lower().startswith("heading") or (len(text) < 110 and text[:1].isdigit() and "." in text[:7])


def render_document(document: Document) -> tuple[str, str]:
    title = ""
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if not title:
            title = text
            continue
        tag = "h2" if is_heading(paragraph.style.name, text) else "p"
        parts.append(f"<{tag}>{escape(text)}</{tag}>")
    for table in document.tables:
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = []
            for cell in row.cells:
                value = " ".join(cell.text.split())
                cell_tag = "th" if row_index == 0 else "td"
                cells.append(f"<{cell_tag}>{escape(value)}</{cell_tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        parts.append("<table>" + "".join(rows) + "</table>")
    return title, "\n".join(parts)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for source, slug in SOURCES:
        document = Document(source)
        title, body = render_document(document)
        (OUTPUT / f"{slug}.html").write_text(page(title, body), encoding="utf-8")


if __name__ == "__main__":
    main()
